"""
AWS Lambda function for Transaction Recording Agent
Integrates with ChatGPT via OpenAI Assistant API
"""

import json
import boto3
import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import io
import uuid

# Initialize AWS S3 client
s3_client = boto3.client('s3')

# Configuration
BUCKET_NAME = os.getenv('S3_BUCKET_NAME', 'your-bucket-name')
EXCEL_FILE_NAME = 'transactions.xlsx'

def create_new_excel():
    """Create a new Excel workbook with headers"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Transactions'
    
    headers = ['Transaction ID', 'Date', 'Beneficiary', 'Amount', 'Purpose', 'Created At']
    ws.append(headers)
    
    # Style headers
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Set column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 20
    
    return wb

def get_excel_from_s3():
    """Fetch existing Excel file from S3 or create new one"""
    try:
        response = s3_client.get_object(Bucket=BUCKET_NAME, Key=EXCEL_FILE_NAME)
        excel_data = response['Body'].read()
        wb = load_workbook(io.BytesIO(excel_data))
        return wb
    except s3_client.exceptions.NoSuchKey:
        # File doesn't exist, create new one
        return create_new_excel()

def save_excel_to_s3(wb):
    """Save Excel workbook to S3"""
    try:
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        s3_client.put_object(
            Bucket=BUCKET_NAME,
            Key=EXCEL_FILE_NAME,
            Body=excel_buffer.getvalue(),
            ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        return True
    except Exception as e:
        print(f"Error saving to S3: {str(e)}")
        return False

def record_transaction(beneficiary: str, date: str, amount: float, purpose: str):
    """
    Record a transaction in Excel file on S3
    
    Args:
        beneficiary: Name of the beneficiary
        date: Transaction date (YYYY-MM-DD)
        amount: Transaction amount
        purpose: Purpose of transaction
    
    Returns:
        dict: Response with transaction ID and status
    """
    try:
        # Validate inputs
        if not all([beneficiary, date, amount, purpose]):
            return {
                'success': False,
                'message': 'Missing required fields: beneficiary, date, amount, purpose'
            }
        
        # Validate date format
        try:
            datetime.strptime(date, '%Y-%m-%d')
        except ValueError:
            return {
                'success': False,
                'message': 'Invalid date format. Use YYYY-MM-DD'
            }
        
        # Validate amount
        try:
            amount_float = float(amount)
            if amount_float <= 0:
                raise ValueError('Amount must be positive')
        except (ValueError, TypeError):
            return {
                'success': False,
                'message': 'Invalid amount. Must be a positive number'
            }
        
        # Get or create Excel file
        wb = get_excel_from_s3()
        ws = wb.active
        
        # Generate transaction ID
        transaction_id = str(uuid.uuid4())[:8].upper()
        
        # Add transaction record
        ws.append([
            transaction_id,
            date,
            beneficiary,
            amount_float,
            purpose,
            datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        ])
        
        # Save back to S3
        if save_excel_to_s3(wb):
            return {
                'success': True,
                'message': f'Transaction recorded successfully',
                'transaction_id': transaction_id,
                'details': {
                    'beneficiary': beneficiary,
                    'date': date,
                    'amount': amount_float,
                    'purpose': purpose
                }
            }
        else:
            return {
                'success': False,
                'message': 'Failed to save transaction'
            }
    
    except Exception as e:
        return {
            'success': False,
            'message': f'Error: {str(e)}'
        }

def lambda_handler(event, context):
    """AWS Lambda handler function"""
    try:
        body = json.loads(event.get('body', '{}'))
        
        action = body.get('action')
        
        if action == 'record_transaction':
            result = record_transaction(
                beneficiary=body.get('beneficiary'),
                date=body.get('date'),
                amount=body.get('amount'),
                purpose=body.get('purpose')
            )
            
            return {
                'statusCode': 200 if result['success'] else 400,
                'body': json.dumps(result)
            }
        
        else:
            return {
                'statusCode': 400,
                'body': json.dumps({'error': 'Invalid action'})
            }
    
    except Exception as e:
        return {
            'statusCode': 500,
            'body': json.dumps({'error': str(e)})
        }
